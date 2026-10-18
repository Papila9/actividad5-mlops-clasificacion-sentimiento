"""Build the technical report from validated project artifacts."""
from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reportes" / "reporte_tecnico.docx"


def load_json(path: str):
    return json.loads((ROOT / path).read_text(encoding="utf-8"))


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = str(value)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def add_image(doc, relative_path, caption, width=6.2):
    path = ROOT / relative_path
    if path.exists():
        doc.add_picture(str(path), width=Inches(width))
        doc.add_paragraph(caption, style="Caption")


def main():
    original = load_json("reportes/entrenamiento/metricas_experimentos.json")
    grouped = load_json("reportes/auditoria_fase3/evaluacion_agrupada/evaluacion_agrupada.json")
    audit = load_json("reportes/auditoria_fase3/auditoria.json")
    manifest = load_json("datos/datos_limp/manifiesto.json")
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.65)
    sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.75)
    sec.right_margin = Inches(.75)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Title"].font.color.rgb = RGBColor(31, 78, 121)

    doc.add_heading("Actividad 5", 0)
    doc.add_paragraph("Preparación, entrenamiento y seguimiento de modelos de clasificación de sentimiento")
    doc.add_paragraph("Informe técnico")
    doc.add_paragraph("Asignatura: _________________________________________________")
    doc.add_paragraph("Estudiante: ______________________________    Grupo: __________")
    doc.add_paragraph("Docente: __________________________________    Institución: _____")
    doc.add_paragraph("Fecha: _____________________________________________________")
    doc.add_paragraph("Proyecto académico reproducible. No es un sistema de producción.")

    doc.add_heading("Resumen", level=1)
    doc.add_paragraph(
        "Se preparó un conjunto de 14,640 comentarios en inglés asociados a aerolíneas, "
        "se compararon Regresión Logística y Random Forest con TF-IDF, y se registraron "
        "cuatro ejecuciones en MLflow: dos modelos bajo partición estratificada aleatoria "
        "y dos bajo evaluación agrupada por texto idéntico. La evaluación agrupada es un "
        "escenario separado y no reemplaza ni se promedia con las métricas originales."
    )
    doc.add_heading("Objetivo y datos", level=1)
    doc.add_paragraph(
        "El objetivo es clasificar cada comentario como negative, neutral o positive usando "
        "el texto (text) y la etiqueta airline_sentiment. El dataset utilizado es Twitter US "
        "Airline Sentiment, publicado en Kaggle: "
        "https://www.kaggle.com/datasets/crowdflower/twitter-airline-sentiment/data. "
        "La página atribuye la fuente original a Crowdflower's Data for Everyone library, "
        "acredita a Figure Eight y un colaborador, y muestra licencia CC BY-NC-SA 4.0 "
        "(consulta registrada: 17-10-2026). No existe registro local de descarga ni checksum "
        "del editor; la procedencia byte a byte del CSV no puede autenticarse."
    )
    doc.add_paragraph(
        "El CSV original tiene 14,640 filas y 15 columnas. La preparación retuvo las dos "
        "variables del modelado, recortó espacios, normalizó las etiquetas a minúsculas y "
        "excluyó textos vacíos o etiquetas inválidas. En los datos observados no se excluyeron "
        "filas por estas dos reglas; se conservaron las repeticiones texto-etiqueta para auditoría."
    )
    doc.add_heading("Diccionario de variables", level=2)
    add_table(doc, ["Variable", "Descripción / uso"], [
        ["tweet_id", "Identificador del tuit; no se utiliza en el modelo."],
        ["airline_sentiment", "Etiqueta objetivo: negative, neutral o positive."],
        ["airline_sentiment_confidence", "Confianza de la anotación; excluida del modelo."],
        ["negativereason", "Motivo anotado para sentimiento negativo; excluido."],
        ["negativereason_confidence", "Confianza del motivo negativo; excluida."],
        ["airline", "Aerolínea mencionada; excluida del modelo."],
        ["airline_sentiment_gold", "Etiqueta suplementaria de referencia; excluida."],
        ["name", "Nombre de cuenta; dato potencialmente identificable, excluido."],
        ["negativereason_gold", "Motivo suplementario de referencia; excluido."],
        ["retweet_count", "Conteo de retuits; excluido."],
        ["text", "Texto del comentario; única entrada del clasificador."],
        ["tweet_coord", "Coordenadas declaradas; excluidas."],
        ["tweet_created", "Fecha del tuit; excluida."],
        ["tweet_location", "Ubicación declarada; excluida."],
        ["user_timezone", "Zona horaria declarada; excluida."],
    ])
    quality = load_json("datos/datos_limp/calidad_datos.json")
    before = quality["before_cleaning"]
    after = quality["after_cleaning"]
    doc.add_heading("Exploración y calidad de los datos", level=2)
    doc.add_paragraph(
        "La fuente contiene valores ausentes en campos que no se conservaron: "
        f"negativereason ({before['null_or_empty_by_column']['negativereason']}), "
        f"negativereason_confidence ({before['null_or_empty_by_column']['negativereason_confidence']}), "
        f"airline_sentiment_gold ({before['null_or_empty_by_column']['airline_sentiment_gold']}), "
        f"negativereason_gold ({before['null_or_empty_by_column']['negativereason_gold']}), "
        f"tweet_coord ({before['null_or_empty_by_column']['tweet_coord']}), "
        f"tweet_location ({before['null_or_empty_by_column']['tweet_location']}) y "
        f"user_timezone ({before['null_or_empty_by_column']['user_timezone']}). "
        f"En las columnas usadas quedaron {after['null_or_empty_by_column']['text']} nulos/vacíos en text y "
        f"{after['null_or_empty_by_column']['airline_sentiment']} en airline_sentiment. Se registraron "
        f"{before['duplicate_full_rows']} filas completas duplicadas en el origen y "
        f"{before['duplicate_text_label_rows']} repeticiones texto-etiqueta; no se eliminaron."
    )
    doc.add_heading("Distribución de clases", level=2)
    total = sum(after["class_counts"].values())
    add_table(doc, ["Etiqueta", "Registros", "Porcentaje"], [
        [label, count, f"{100*count/total:.2f}%"] for label, count in after["class_counts"].items()
    ])
    class_plot = ROOT / "evidencias" / "distribucion_clases_reporte.png"
    fig, ax = plt.subplots(figsize=(6, 3.2))
    labels, counts = list(after["class_counts"].keys()), list(after["class_counts"].values())
    ax.bar(labels, counts, color=["#4472C4", "#ED7D31", "#70AD47"])
    ax.set(title="Distribución de clases", ylabel="Registros")
    fig.tight_layout()
    fig.savefig(class_plot, dpi=160)
    plt.close(fig)
    add_image(doc, "evidencias/distribucion_clases_reporte.png", "Figura. Distribución de las clases en el dataset limpio.")
    add_table(doc, ["Artefacto", "Filas", "SHA-256"], [
        ["CSV de origen local", manifest["source"]["rows"], manifest["source"]["sha256"]],
        ["CSV limpio", manifest["processed"]["rows"], manifest["processed"]["sha256"]],
    ])

    doc.add_heading("Método y control de evaluación", level=1)
    doc.add_paragraph(
        "El escenario original utiliza partición estratificada 80/20 con semilla 42 "
        "(11,712 entrenamiento y 2,928 prueba). GridSearchCV seleccionó por F1 macro con "
        "cinco folds estratificados. La transformación TF-IDF forma parte del pipeline y se "
        "ajusta dentro de cada fold únicamente con sus datos de entrenamiento. El conjunto "
        "de prueba no se usa para seleccionar hiperparámetros."
    )
    doc.add_paragraph(
        "La auditoría encontró 71 textos únicos compartidos entre entrenamiento y prueba: "
        "60 con etiqueta consistente en esos conjuntos y 11 con etiquetas distintas. En todo "
        "el dataset existen 25 grupos de texto con etiquetas conflictivas. Como evaluación "
        "adicional se usó StratifiedGroupKFold (5 folds, semilla 42), reservando el primer fold "
        "para prueba y agrupando también la validación interna. En este protocolo no hubo "
        "solapamiento de textos idénticos entre entrenamiento y prueba."
    )
    doc.add_heading("Modelos y búsqueda de hiperparámetros", level=2)
    doc.add_paragraph(
        "Regresión Logística es un clasificador lineal multinomial aplicado a los vectores dispersos "
        "TF-IDF. Random Forest combina árboles de decisión construidos sobre subconjuntos de variables "
        "y observaciones; se evaluó sobre una representación TF-IDF limitada en vocabulario."
    )
    param_rows = []
    for scenario, source in (("Original", original["models"]), ("Agrupada", grouped["models"])):
        for name, model in source.items():
            cvmean = model.get("cv_selected_candidate", {}).get("mean_f1_macro", model.get("cv_mean_f1_macro"))
            cvstd = model.get("cv_selected_candidate", {}).get("std_f1_macro", model.get("cv_std_f1_macro_sample"))
            duration = model.get("fit_duration_seconds")
            param_rows.append([scenario, name, json.dumps(model["best_params"], ensure_ascii=False),
                               f"{cvmean:.4f} ± {cvstd:.4f}", f"{duration:.3f} s"])
    add_table(doc, ["Protocolo", "Modelo", "Mejores parámetros", "F1 macro CV", "Búsqueda (tiempo)"], param_rows)
    doc.add_paragraph(
        "La búsqueda comparó cuatro candidatos por modelo y escenario. En el protocolo original "
        "el criterio de selección fue F1 macro. Las métricas de prueba se calcularon después de "
        "seleccionar los parámetros; no se usaron para ajustarlos. Los tiempos corresponden a "
        "GridSearchCV registrado en los artefactos de cada escenario."
    )
    doc.add_heading("Resultados", level=1)
    rows = []
    for scenario, source in (("Original estratificada", original["models"]), ("Agrupada por texto", grouped["models"])):
        for name, model in source.items():
            m = model["test_metrics"]
            cvmean = model.get("cv_selected_candidate", {}).get("mean_f1_macro", model.get("cv_mean_f1_macro"))
            cvstd = model.get("cv_selected_candidate", {}).get("std_f1_macro", model.get("cv_std_f1_macro_sample"))
            rows.append([scenario, "Regresión Logística" if name == "logistic_regression" else "Random Forest",
                         f"{m['accuracy']:.4f}", f"{m['precision_macro']:.4f}", f"{m['recall_macro']:.4f}",
                         f"{m['f1_macro']:.4f}", f"{m['f1_weighted']:.4f}", f"{m['roc_auc_ovr_macro']:.4f}",
                         f"{cvmean:.4f} ± {cvstd:.4f}"])
    add_table(doc, ["Escenario", "Modelo", "Accuracy", "Prec. macro", "Recall macro", "F1 macro", "F1 ponderado", "ROC-AUC OVR", "F1 CV media ± DE"], rows)
    doc.add_paragraph(
        "Las métricas de prueba y las matrices/reportes proceden de los artefactos validados. "
        "La reproducción de la auditoría confirmó parámetros, folds, métricas y conteos de "
        "confusión en ambos escenarios. La selección se interpreta descriptivamente: no se "
        "realizó una prueba de significancia estadística y se usó una sola semilla."
    )
    for scenario, base in (("Original", "reportes/entrenamiento"),
                           ("Agrupado", "reportes/auditoria_fase3/evaluacion_agrupada")):
        doc.add_heading(f"Reporte por clase — {scenario}", level=2)
        rows = []
        for slug, label in (("logistic_regression", "Regresión Logística"), ("random_forest", "Random Forest")):
            report = load_json(f"{base}/classification_report_{slug}.json")
            for klass in ("negative", "neutral", "positive"):
                entry = report[klass]
                rows.append([label, klass, f"{entry['precision']:.3f}", f"{entry['recall']:.3f}",
                             f"{entry['f1-score']:.3f}", int(entry["support"])])
        add_table(doc, ["Modelo", "Clase", "Precisión", "Recall", "F1", "Soporte"], rows)
    add_image(doc, "evidencias/comparacion_modelos.png", "Figura 1. Comparación original de modelos.")
    add_image(doc, "evidencias/matrices_confusion.png", "Figura 2. Matrices de confusión del escenario original.")
    add_image(doc, "evidencias/auditoria_fase3/evaluacion_agrupada/comparacion_modelos.png", "Figura 3. Comparación del escenario agrupado.")
    add_image(doc, "evidencias/auditoria_fase3/evaluacion_agrupada/matrices_confusion.png", "Figura 4. Matrices del escenario agrupado.")

    doc.add_heading("Seguimiento con MLflow", level=1)
    doc.add_paragraph(
        "El experimento local Actividad5_Clasificacion_Sentimiento contiene cuatro ejecuciones "
        "confirmadas en la interfaz (MLflow 3.17.1): logistic regression y random forest para "
        "evaluacion_original y evaluacion_agrupada. Los runs registran métricas, parámetros, "
        "versiones y artefactos; sus etiquetas marcan el escenario y especifican que la prueba "
        "no se usó en ajuste. Las ejecuciones registradas no guardan contenido de los comentarios "
        "ni serializan modelos. Las capturas insertadas son evidencia visual de la interfaz; "
        "la base SQLite y el almacenamiento interno de artefactos quedan fuera de este informe/Git."
    )
    for p in sorted((ROOT / "evidencias" / "mlflow").glob("*.png")):
        doc.add_picture(str(p), width=Inches(6.2))
        doc.add_paragraph(f"Evidencia MLflow: {p.name}", style="Caption")

    doc.add_heading("Pruebas y reproducibilidad", level=1)
    doc.add_paragraph(
        "La suite unittest disponible reportó 7 pruebas OK en la última ejecución. "
        "pip check no reportó conflictos. Las dependencias se instalaron en el entorno virtual; "
        "el notebook complementario se validó como JSON y compilando estáticamente sus celdas, "
        "pero no se ejecutó integralmente porque no están instalados Jupyter ni ipykernel."
    )
    doc.add_paragraph(
        "Para reproducir: crear .venv con Python 3.13, instalar requirements.txt, descargar "
        "Tweets.csv desde la URL de Kaggle y guardarlo en datos/datos_ini/, ejecutar "
        "fuentes/datos_prep.py, fuentes/train.py, fuentes/auditoria_fase3.py, "
        "fuentes/mlflow_tracking.py y después las pruebas con "
        ".venv/Scripts/python.exe -m unittest discover -s tests -v. Para consultar MLflow, "
        "ejecutar .venv/Scripts/mlflow.exe server --backend-store-uri sqlite:///mlflow.db "
        "--host 127.0.0.1 --port 5000 y abrir http://127.0.0.1:5000."
    )
    doc.add_heading("Limitaciones y decisiones", level=1)
    doc.add_paragraph(
        "Los textos duplicados pueden influir en la estimación del split aleatorio; por eso se "
        "reportan dos protocolos separados. Hay desbalance de clases y etiquetas contradictorias "
        "para algunos textos. La licencia de Kaggle es no comercial y share-alike; por cautela "
        "el CSV no se incluye en Git y se documenta su descarga. No se evaluó desempeño en "
        "producción ni generalización a otros dominios."
    )
    doc.add_paragraph(
        "Deuda técnica y trabajo futuro: repetir la evaluación con más semillas o validación "
        "externa, revisar el efecto de etiquetas contradictorias, cuantificar intervalos de "
        "incertidumbre y sesgos por aerolínea, y habilitar una ejecución completa del notebook "
        "en un entorno con Jupyter. Estas son propuestas; no se reportan como experimentos realizados."
    )
    doc.add_heading("Conclusiones", level=1)
    doc.add_paragraph(
        "En las ejecuciones disponibles, Regresión Logística obtuvo mayor F1 macro que Random Forest "
        "en ambos protocolos. Su F1 macro fue 0.7457 en la partición original y 0.7437 en la agrupada; "
        "la cercanía entre esos resultados es descriptiva y no prueba equivalencia ni significancia. "
        "La evaluación agrupada aporta una medición complementaria que elimina el solapamiento de "
        "textos idénticos. El flujo registra datos, resultados y evidencias de manera reproducible, "
        "sujeto a las limitaciones de la fuente y de los protocolos documentados."
    )
    doc.add_heading("Referencias", level=1)
    doc.add_paragraph(
        "Crowdflower. Twitter US Airline Sentiment. Kaggle, dataset y licencia CC BY-NC-SA 4.0. "
        "Consulta: 17 de octubre de 2026. https://www.kaggle.com/datasets/crowdflower/twitter-airline-sentiment/data"
    )
    doc.add_paragraph(
        "Scikit-learn developers. Pipeline y composición de estimadores. "
        "https://scikit-learn.org/stable/modules/compose.html"
    )
    doc.add_paragraph(
        "Scikit-learn developers. GridSearchCV. https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.GridSearchCV.html"
    )
    doc.add_paragraph("MLflow. Tracking experiments. https://mlflow.org/docs/latest/ml/tracking/")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"Informe generado: {OUT.relative_to(ROOT)} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
